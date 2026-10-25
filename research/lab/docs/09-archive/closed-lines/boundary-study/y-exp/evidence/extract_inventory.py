from collections import defaultdict
from datetime import datetime
from pathlib import Path
import json

from openpyxl import load_workbook


INPUT = Path('/Users/mr.simi/Desktop/26.秦沈项目/04.成本与资金/02.标后预算/标后预算成本数据/2秦沈项目标后预算明细表-1.15局会议修改后.xlsx')
OUTDIR = Path(__file__).resolve().parent


def cell_text(value):
    return '' if value is None else str(value).replace('\n', '\\n')


wb_formula = load_workbook(INPUT, data_only=False, read_only=False, keep_links=True)
wb_cached = load_workbook(INPUT, data_only=True, read_only=True, keep_links=True)

inventory = {
    'input': str(INPUT),
    'extracted_at': datetime.now().isoformat(timespec='seconds'),
    'sheets': {},
    'defined_names': [],
    'external_links': [],
}

for prop in ('title', 'creator', 'lastModifiedBy', 'created', 'modified'):
    inventory[prop] = str(getattr(wb_formula.properties, prop, None))

for defined_name in wb_formula.defined_names.values():
    inventory['defined_names'].append({
        'name': defined_name.name,
        'value': defined_name.value,
    })

try:
    inventory['external_links'] = [link.file_link.target if link.file_link else str(link) for link in wb_formula._external_links]
except Exception as exc:
    inventory['external_links'] = [f'读取失败: {exc!r}']

for title in wb_formula.sheetnames:
    ws = wb_formula[title]
    wsc = wb_cached[title]
    matches = []
    comments = []
    validations = []
    e_formulas = []
    e_values = []

    for row in ws.iter_rows():
        for cell in row:
            text = cell_text(cell.value)
            if text and any(keyword in text for keyword in ('万元', '单位', '金额', '单价', '元/', '合计')):
                matches.append({
                    'sheet': title,
                    'cell': cell.coordinate,
                    'text': text,
                })
            if cell.comment is not None:
                comments.append({
                    'cell': cell.coordinate,
                    'author': cell.comment.author,
                    'text': cell_text(cell.comment.text),
                })

    for rule in ws.data_validations.dataValidation:
        validations.append({
            'ranges': str(rule.sqref),
            'type': rule.type,
            'operator': rule.operator,
            'formula1': rule.formula1,
            'formula2': rule.formula2,
            'prompt': rule.prompt,
            'promptTitle': rule.promptTitle,
            'error': rule.error,
            'errorTitle': rule.errorTitle,
        })

    for row in range(1, min(ws.max_row, 200) + 1):
        formula_cell = ws.cell(row, 5)
        cached_cell = wsc.cell(row, 5)
        if formula_cell.value is not None or cached_cell.value is not None:
            e_formulas.append({
                'row': row,
                'formula': cell_text(formula_cell.value),
                'cached': cached_cell.value,
                'A': cell_text(ws.cell(row, 1).value),
                'B': cell_text(ws.cell(row, 2).value),
                'C': cell_text(ws.cell(row, 3).value),
                'D': cell_text(ws.cell(row, 4).value),
                'F': cell_text(ws.cell(row, 6).value),
                'G': cell_text(ws.cell(row, 7).value),
                'H': cell_text(ws.cell(row, 8).value),
            })

    header_rows = []
    for row in range(1, min(ws.max_row, 12) + 1):
        row_values = []
        for col in range(1, min(ws.max_column, 20) + 1):
            row_values.append(cell_text(ws.cell(row, col).value))
        if any(row_values):
            header_rows.append({'row': row, 'values': row_values})

    inventory['sheets'][title] = {
        'state': ws.sheet_state,
        'max_row': ws.max_row,
        'max_column': ws.max_column,
        'merged': [str(r) for r in ws.merged_cells.ranges],
        'freeze_panes': ws.freeze_panes,
        'keyword_matches': matches,
        'comments': comments,
        'data_validations': validations,
        'e_nonempty_rows_to_200': e_formulas,
        'header_candidate_rows': header_rows,
    }

inventory['sheets'][title]['cached_max_row'] = wsc.max_row
inventory['sheets'][title]['cached_max_column'] = wsc.max_column

(OUTDIR / 'inventory.json').write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding='utf-8')

print('工作表：')
for name, meta in inventory['sheets'].items():
    print(f"- {name}: state={meta['state']}, size={meta['max_row']}x{meta['max_column']}, merged={len(meta['merged'])}, comments={len(meta['comments'])}, validations={len(meta['data_validations'])}, E非空(前200)={len(meta['e_nonempty_rows_to_200'])}, 单位/金额关键词={len(meta['keyword_matches'])}")
print(f"定义名称：{len(inventory['defined_names'])}；外部链接：{len(inventory['external_links'])}")
print('输出：inventory.json')
