import ast
import csv
import hashlib
import os
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


INPUT = Path("/Users/mr.simi/Desktop/26.秦沈项目/04.成本与资金/02.标后预算/标后预算成本数据/2秦沈项目标后预算明细表-1.15局会议修改后.xlsx")
BACKUP = INPUT.with_name("2秦沈项目标后预算明细表-1.15局会议修改后_D列替换前备份_20261002.xlsx")
TEMP = INPUT.with_name(INPUT.name + ".tmp-d-unit-replace.xlsx")
OLD_SHA = "1dd8372a194135049137b2112431043fcb4e948c0532a3c05272d7a136adef51"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shared_strings(data):
    root = ET.fromstring(data)
    return ["".join(node.text or "" for node in si.iter(f"{{{NS['m']}}}t")) for si in root.findall("m:si", NS)]


def cells(worksheet_data):
    root = ET.fromstring(worksheet_data)
    return list(root.iter(f"{{{NS['m']}}}c"))


assert sha256(INPUT) == OLD_SHA
if not BACKUP.exists():
    shutil.copy2(INPUT, BACKUP)
assert sha256(BACKUP) == OLD_SHA

with Path("d_unit_audit.tsv").open(encoding="utf-8") as source:
    targets = list(csv.DictReader(source, delimiter="\t"))
assert len(targets) == 358 and len({row["row"] for row in targets}) == 358

with zipfile.ZipFile(INPUT) as archive:
    infos = archive.infolist()
    parts = {info.filename: archive.read(info.filename) for info in infos}

old_shared = shared_strings(parts["xl/sharedStrings.xml"])
suggested_to_index = {text: index for index, text in reversed(list(enumerate(old_shared)))}
new_shared = []
for target in targets:
    suggested = target["suggested"]
    if suggested not in suggested_to_index:
        old_shared.append(suggested)
        new_shared.append(f"<si><t>{suggested}</t></si>")
        suggested_to_index[suggested] = len(old_shared) - 1

shared_text = parts["xl/sharedStrings.xml"].decode("utf-8")
assert shared_text.count("</sst>") == 1
shared_text = shared_text.replace("</sst>", "".join(new_shared) + "</sst>")
assert shared_text.count('uniqueCount="1864"') == 1
shared_text = shared_text.replace(f'uniqueCount="{1864}"', f'uniqueCount="{len(old_shared)}"')
parts["xl/sharedStrings.xml"] = shared_text.encode("utf-8")

sheet_name = "xl/worksheets/sheet2.xml"
original_sheet_data = parts[sheet_name]
sheet_text = parts[sheet_name].decode("utf-8")
changed = 0
for target in targets:
    row_no = int(target["row"])
    match = re.search(rf'<c r="D{row_no}"(?:\s+[^>]*)?>.*?</c>', sheet_text)
    assert match
    old_cell = match.group(0)
    value_match = re.search(r"<v>(\d+)</v>", old_cell)
    assert value_match and 't="s"' in old_cell
    assert old_shared[int(value_match.group(1))] == ast.literal_eval(target["D_raw_repr"])
    new_cell = old_cell.replace(value_match.group(0), f"<v>{suggested_to_index[target['suggested']]}</v>")
    sheet_text = sheet_text.replace(old_cell, new_cell, 1)
    changed += 1
assert changed == 358
parts[sheet_name] = sheet_text.encode("utf-8")

with zipfile.ZipFile(INPUT) as source_archive:
    changed_parts = {info.filename for info in source_archive.infolist() if parts[info.filename] != source_archive.read(info.filename)}
assert changed_parts == {sheet_name, "xl/sharedStrings.xml"}

with zipfile.ZipFile(TEMP, "w") as output_archive:
    for info in infos:
        output_archive.writestr(info, parts[info.filename])

old_cells = cells(original_sheet_data)
new_cells = cells(sheet_text.encode("utf-8"))
target_rows = {int(target["row"]) for target in targets}
assert len(old_cells) == len(new_cells)
for old_cell, new_cell in zip(old_cells, new_cells):
    reference = new_cell.attrib["r"]
    assert old_cell.attrib == new_cell.attrib
    assert len(old_cell) == len(new_cell)
    for old_child, new_child in zip(old_cell, new_cell):
        assert old_child.tag == new_child.tag and old_child.attrib == new_child.attrib
        value_change = reference.lstrip("D") in {str(row) for row in target_rows} and reference.startswith("D") and old_child.tag.endswith("}v")
        if not value_change:
            assert old_child.text == new_child.text

old_workbook = load_workbook(BACKUP, data_only=False)
new_workbook = load_workbook(TEMP, data_only=False)
assert old_workbook.sheetnames == new_workbook.sheetnames
target_suggested = {int(target["row"]): target["suggested"] for target in targets}
for sheet_title in old_workbook.sheetnames:
    old_sheet = old_workbook[sheet_title]
    new_sheet = new_workbook[sheet_title]
    assert old_sheet.max_row == new_sheet.max_row and old_sheet.max_column == new_sheet.max_column
    assert old_sheet.merged_cells.ranges == new_sheet.merged_cells.ranges
    assert old_sheet.freeze_panes == new_sheet.freeze_panes
    for row_no, (old_row, new_row) in enumerate(zip(old_sheet.iter_rows(values_only=True), new_sheet.iter_rows(values_only=True)), 1):
        for col_no, (old_value, new_value) in enumerate(zip(old_row, new_row), 1):
            expected_change = sheet_title == "表2-分部分项清单" and col_no == 4 and row_no in target_suggested
            if expected_change:
                assert new_value == target_suggested[row_no]
            else:
                assert old_value == new_value
old_workbook.close()
new_workbook.close()

check_workbook = load_workbook(TEMP, read_only=True, data_only=True)
check_sheet = check_workbook["表2-分部分项清单"]
remaining = []
for row in check_sheet.iter_rows(min_row=5, values_only=True):
    if row[3] is not None and str(row[3]).strip():
        remaining.append(str(row[3]).strip())
check_workbook.close()
assert len(remaining) == 841 and len(set(remaining)) == 11
assert set(remaining) == {"m³", "kg", "m", "m²", "个", "总额", "套", "座", "块", "dm³", "t"}

os.replace(TEMP, INPUT)
assert sha256(BACKUP) == OLD_SHA
assert sha256(INPUT) != OLD_SHA
print(f"backup={BACKUP}")
print(f"old_sha256={OLD_SHA}")
print(f"new_sha256={sha256(INPUT)}")
print(f"changed_cells={changed}")
print(f"changed_parts={sorted(changed_parts)}")
print(f"remaining_unique_d={len(set(remaining))}")
