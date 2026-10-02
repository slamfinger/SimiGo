from pathlib import Path

from openpyxl import load_workbook


BASE = Path('/Users/mr.simi/Desktop/26.秦沈项目/04.成本与资金/02.标后预算/标后预算成本数据')
FILES = [
    '2秦沈项目标后预算明细表-1.15公司成本部上报.xlsx',
    '2秦沈项目标后预算明细表-1.15局会议修改后 - 梁成刚版.xlsx',
    '2秦沈项目标后预算明细表-1.15局会议修改后_复核.xlsx',
    '秦沈项目标后表-12.28.xlsx',
]
TARGET = '表2-分部分项清单'
KEYS = [('-a', '挖土方', '挖土方'), ('-d', '借土填方', '借土填方')]
OUT = Path(__file__).resolve().parent / 'versions_compare.tsv'


def text(v):
    return '' if v is None else str(v).replace('\t', ' ').replace('\n', ' ')


rows = [['file', 'sheet', 'D2', 'E2', 'row', 'A', 'B', 'D', 'E', 'row', 'A', 'B', 'D', 'E']]
for file_name in FILES:
    path = BASE / file_name
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = TARGET if TARGET in wb.sheetnames else next((s for s in wb.sheetnames if '分部分项' in s), None)
    if not sheet_name:
        rows.append([file_name, '<无分部分项工作表>', '', '', '', '', '', ''])
        continue
    ws = wb[sheet_name]
    header_d = ws['D2'].value
    header_e = ws['E2'].value
    rows.append([file_name, sheet_name, text(header_d), text(header_e)])
    for a, b, expected in KEYS:
        found = ''
        for row in ws.iter_rows(min_row=3, max_col=8):
            vals = [text(c.value) for c in row]
            if len(vals) >= 5 and vals[0] == a and vals[1] == b and expected in vals[1] + vals[2]:
                found = '\t'.join([str(row[0].row), a, b, vals[3], vals[4]])
                break
        rows[-1].extend(found.split('\t') if found else [''] * 5)
    wb.close()

OUT.write_text('\n'.join('\t'.join(r) for r in rows) + '\n', encoding='utf-8')
print(OUT.read_text(encoding='utf-8'))
