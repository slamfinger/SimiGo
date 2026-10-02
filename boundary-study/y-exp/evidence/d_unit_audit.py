import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


INPUT = Path("/Users/mr.simi/Desktop/26.秦沈项目/04.成本与资金/02.标后预算/标后预算成本数据/2秦沈项目标后预算明细表-1.15局会议修改后.xlsx")
CANON = {"m3": "m³", "m2": "m²", "㎡": "m²", "dm3": "dm³", "kg": "kg"}


wb = load_workbook(INPUT, read_only=True, data_only=True)
ws = wb["表2-分部分项清单"]
rows = []
both_count = 0
d_only_rows = []
e_only_rows = []

for row_no, row in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
    raw = row[3] if len(row) > 3 else None
    e_cached = row[4] if len(row) > 4 else None
    has_e = e_cached is not None and str(e_cached).strip() != ""
    if raw is not None and str(raw).strip():
        value = str(raw)
        trimmed = value.strip()
        rows.append({
            "row": row_no,
            "raw_repr": repr(value),
            "trimmed": trimmed,
            "suggested": CANON.get(trimmed, trimmed),
            "e_cached": e_cached,
        })
        if has_e:
            both_count += 1
        else:
            d_only_rows.append(row_no)
    elif has_e:
        e_only_rows.append(row_no)

wb.close()

by_raw = defaultdict(list)
for record in rows:
    by_raw[record["raw_repr"]].append(record["row"])

changes = [record for record in rows if record["raw_repr"] != repr(record["suggested"])]

with Path("d_unit_audit.tsv").open("w", encoding="utf-8") as f:
    f.write("row\tD_raw_repr\tD_trimmed\tsuggested\tE_cached\n")
    for record in changes:
        f.write(f'{record["row"]}\t{record["raw_repr"]}\t{record["trimmed"]}\t{record["suggested"]}\t{record["e_cached"]}\n')

summary = {
    "input": str(INPUT),
    "sheet": "表2-分部分项清单",
    "audited_min_row": 5,
    "audited_nonempty_d": len(rows),
    "coverage": {
        "d_and_e": both_count,
        "d_only": len(d_only_rows),
        "e_only": len(e_only_rows),
        "d_only_rows": d_only_rows,
        "e_only_rows": e_only_rows,
    },
    "unique_raw_d": len(by_raw),
    "raw_counts": dict(Counter(record["raw_repr"] for record in rows).most_common()),
    "suggested_counts": dict(Counter(record["suggested"] for record in rows).most_common()),
    "change_rows": len(changes),
    "changes_by_rule": {f"{raw} -> {suggested}": count for (raw, suggested), count in Counter(
        (record["raw_repr"], record["suggested"]) for record in changes
    ).items()},
    "rows_by_raw": {raw: row_numbers for raw, row_numbers in sorted(by_raw.items())},
}
Path("d_unit_audit.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
