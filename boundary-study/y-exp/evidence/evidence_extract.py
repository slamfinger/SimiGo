import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


INPUT = Path('/Users/mr.simi/Desktop/26.秦沈项目/04.成本与资金/02.标后预算/标后预算成本数据/2秦沈项目标后预算明细表-1.15局会议修改后.xlsx')
OUTDIR = Path(__file__).resolve().parent
SHEET = '表2-分部分项清单'
E_REF = re.compile(r'(?<![A-Za-z$])E(\d+)(?![0-9])')


def clean(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return f'{v:.10g}'
    return str(v).replace('\t', ' ').replace('\n', '\\n')


formula_wb = load_workbook(INPUT, data_only=False, read_only=True)
cached_wb = load_workbook(INPUT, data_only=True, read_only=True)
wsf = formula_wb[SHEET]
wsv = cached_wb[SHEET]

records = []
for row in wsf.iter_rows(min_row=1, max_row=wsf.max_row, max_col=min(wsf.max_column, 21)):
    row_num = row[0].row
    fs = {get_column_letter(col): cell.value for col, cell in enumerate(row, 1)}
    vs = {get_column_letter(col): cell.value for col, cell in enumerate(wsv.iter_rows(min_row=row_num, max_row=row_num, max_col=21).__next__(), 1)}
    records.append((row_num, fs, vs))

e_rows = [(n, f, v) for n, f, v in records if f.get('E') not in (None, '')]
numeric_e = [r for r in e_rows if isinstance(r[2]['E'], (int, float))]
formula_e = [r for r in e_rows if isinstance(r[1]['E'], str) and r[1]['E'].startswith('=')]
unit_counts = Counter(clean(r[2].get('D')) for r in e_rows)

e_reference_usage = []
for n, f, v in records:
    for col, val in f.items():
        if isinstance(val, str) and val.startswith('=') and E_REF.search(val):
            e_reference_usage.append({
                'cell': f'{col}{n}',
                'formula': val,
                'cached': v.get(col),
            })

selected_numbers = [2, 5, 12, 35, 44, 47, 54, 75, 76, 77, 78, 79, 112, 184, 185, 186, 187, 188, 203]
selected = [r for r in records if r[0] in selected_numbers]

with (OUTDIR / 'evidence_selected.tsv').open('w', encoding='utf-8') as out:
    out.write('\t'.join(['row', 'A', 'B', 'D', 'E_formula', 'E_cached', 'G_formula', 'G_cached', 'J_formula', 'J_cached', 'R']) + '\n')
    for n, f, v in selected:
        out.write('\t'.join(map(clean, [n, f.get('A'), f.get('B'), f.get('D'), f.get('E'), v.get('E'), f.get('G'), v.get('G'), f.get('J'), v.get('J'), f.get('R')])) + '\n')

with (OUTDIR / 'e_reference_usage.tsv').open('w', encoding='utf-8') as out:
    out.write('cell\tformula\tcached\n')
    for item in e_reference_usage:
        out.write(f"{item['cell']}\t{clean(item['formula'])}\t{clean(item['cached'])}\n")

summary = {
    'sheet': SHEET,
    'header': {
        'D2:D4': wsv['D2'].value,
        'E2:E4': wsv['E2'].value,
        'I2:K2': wsv['I2'].value,
        'I3:I4': wsv['I3'].value,
        'J3:J4': wsv['J3'].value,
    },
    'nonempty_E_rows': len(e_rows),
    'cached_numeric_E_rows': len(numeric_e),
    'hardcoded_numeric_E_rows': len(numeric_e) - len(formula_e),
    'formula_E_rows': len(formula_e),
    'unit_counts': dict(unit_counts),
    'E_min_max': {
        'min': min(v['E'] for _, _, v in numeric_e),
        'max': max(v['E'] for _, _, v in numeric_e),
    },
    'formulas_referencing_E': len(e_reference_usage),
    'reference_formula_columns': dict(Counter(item['cell'].rstrip('0123456789') for item in e_reference_usage)),
}
(OUTDIR / 'e_unit_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')

print(json.dumps(summary, ensure_ascii=False, indent=2))
print('\n选中证据：')
print((OUTDIR / 'evidence_selected.tsv').read_text(encoding='utf-8'))
